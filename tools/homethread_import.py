"""
Load her Homethread.xlsx workbook (and the design files on the flash drive)
straight into the Supabase database.

    python tools/homethread_import.py D:\\Homethread.xlsx --designs "D:\\Embroidery Designs" --dry-run
    python tools/homethread_import.py D:\\Homethread.xlsx --designs "D:\\Embroidery Designs"

Ids are deterministic (sha256 of the row's identity, "imp-" prefix, same scheme
as the app's CSV importer), so running it again updates rows instead of
duplicating them. Sheets read:

  Startup & Fixed Costs  -> expense (is_startup = true), categories created by her names
  Blank Inventory        -> blank, plus an expense per purchase day unless the
                            startup sheet already has that vendor+day+amount
  Product Pricing        -> product (price = Typical Total Price)
  Orders                 -> customer, order, order_line, payment
  Design Library         -> design (+ stitch counts from the drive when a name matches)

Writes go through `npx supabase db query --linked`, one statement per line
(the Windows npx shim can't carry newlines), batched to stay under the
command-line limit.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import subprocess
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from designs_csv import parse_exp, parse_pes, SKIP_DIRS  # noqa: E402

NOW = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")


# ---------------------------------------------------------------------------
# helpers

def imp_id(key: str) -> str:
    return "imp-" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:32]


def cents(v) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, str):
        v = v.replace("$", "").replace(",", "").strip()
        if not v:
            return 0
    try:
        return int(round(float(v) * 100))
    except (TypeError, ValueError):
        return 0


def iso(v):
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str) and re.match(r"^\d{4}-\d{2}-\d{2}$", v.strip()):
        return v.strip()
    return None


def s(v) -> str:
    return "" if v is None else str(v).strip()


def q(v) -> str:
    """SQL literal."""
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''").replace("\n", " ").replace("\r", " ") + "'"


def upsert(table: str, row: dict, pk="id") -> str:
    cols = list(row.keys())
    vals = ", ".join(q(row[c]) for c in cols)
    sets = ", ".join(f"{c} = excluded.{c}" for c in cols if c != pk)
    return f'insert into public."{table}" ({", ".join(cols)}) values ({vals}) on conflict ({pk}) do update set {sets};'


def sheet_rows(ws):
    rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
    hdr_i = max(range(len(rows)), key=lambda i: sum(1 for c in rows[i] if c is not None))
    hdr = [s(c) for c in rows[hdr_i]]
    out = []
    for r in ws.iter_rows(min_row=hdr_i + 2, values_only=True):
        rec = {hdr[i]: r[i] for i in range(min(len(hdr), len(r))) if hdr[i]}
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# sheets

CATEGORY_MAP = {
    "equipment": "Machine & equipment",
    "supplies": "Supplies",
    "design files": "Software & designs",
    "blanks": "Blanks & garments",
}


def load_expenses(ws, cats: dict) -> list[dict]:
    out = []
    for r in sheet_rows(ws):
        d = iso(r.get("Date"))
        amt = cents(r.get("Amount Paid ($)"))
        if not d or amt <= 0:
            continue
        vendor = s(r.get("Vendor")) or "Unknown vendor"
        note = s(r.get("Item / Description"))
        extra = s(r.get("Notes"))
        if extra:
            note = f"{note} ({extra})"
        cat_name = CATEGORY_MAP.get(s(r.get("Category")).lower(), s(r.get("Category")) or "Other")
        kind = s(r.get("Type")).lower()
        out.append({
            "id": imp_id(f"expense|{d}|{vendor.lower()}|{amt}|{note.lower()}"),
            "vendor": vendor,
            "spent_on": d,
            "amount": amt,
            "tax": 0,
            "category_id": cats.setdefault(cat_name, imp_id(f"category|{cat_name.lower()}")),
            "note": note,
            "receipt_image_path": None,
            "extraction_json": None,
            "is_startup": ("startup" in kind) or ("design" in kind),
            "created_at": NOW,
        })
    return out


def load_blanks(ws, cats: dict, expenses: list[dict]) -> tuple[list[dict], list[dict]]:
    blanks, extra_expenses = [], []
    by_day: dict[tuple, list[dict]] = {}
    for r in sheet_rows(ws):
        typ = s(r.get("Blank Type"))
        if not typ:
            continue
        qty = int(r.get("Qty Purchased") or 0)
        total = cents(r.get("Total Cost ($)"))
        d = iso(r.get("Purchase Date"))
        style = s(r.get("Style / Color"))
        vendor = s(r.get("Vendor")) or "Amazon"
        b = {
            "id": imp_id(f"blank|{typ.lower()}|{style.lower()}|{d}|{total}"),
            "type": typ,
            "style": style,
            "vendor": vendor,
            "purchased_on": d,
            "qty": qty,
            "total_cost": total,
            "unit_cost": round(total / qty) if qty else 0,
            "adjust": -int(r.get("Qty Gift/Sample") or 0),
            "notes": s(r.get("Notes")),
            "created_at": NOW,
        }
        blanks.append(b)
        by_day.setdefault((d, vendor.lower()), []).append(b)

    # Blank purchases are money out too. Add an expense per purchase day unless
    # the startup sheet already recorded that vendor+day for about the same money.
    for (d, vendor), group in by_day.items():
        total = sum(b["total_cost"] for b in group)
        already = [e for e in expenses if e["spent_on"] == d and e["vendor"].lower() == vendor]
        if any(abs(e["amount"] - total) <= 150 for e in already) or any(
            abs(e["amount"] - b["total_cost"]) <= 5 and "sash" in e["note"].lower() for e in already for b in group
        ):
            continue
        note = "Blanks: " + "; ".join(f'{b["qty"]} {b["type"]} {b["style"]}'.strip() for b in group)
        extra_expenses.append({
            "id": imp_id(f"expense|{d}|{vendor}|{total}|{note.lower()}"),
            "vendor": group[0]["vendor"],
            "spent_on": d,
            "amount": total,
            "tax": 0,
            "category_id": cats.setdefault("Blanks & garments", "cat-blanks"),
            "note": note,
            "receipt_image_path": None,
            "extraction_json": None,
            "is_startup": True,
            "created_at": NOW,
        })
    return blanks, extra_expenses


def load_products(ws) -> list[dict]:
    out = []
    for i, r in enumerate(sheet_rows(ws)):
        name = s(r.get("Product"))
        if not name:
            continue
        option = s(r.get("Option / Configuration")).replace("\u2013", "-").replace("\ufffd", "-")
        price = cents(r.get("Typical Total Price ($)")) or cents(r.get("Base Selling Price ($)"))
        notes = []
        if cents(r.get("Second-Side / Extra Placement Surcharge ($)")):
            notes.append(f"includes second-side +${cents(r.get('Second-Side / Extra Placement Surcharge ($)'))/100:.0f}")
        out.append({
            "id": imp_id(f"product|{name.lower()}|{option.lower()}"),
            "name": name,
            "option": option,
            "qty": int(r.get("Qty in Option") or 1),
            "price": price,
            "notes": "; ".join(notes),
            "position": i,
            "created_at": NOW,
        })
    return out


def pick_blank(product: str, config: str, blanks: list[dict]):
    """Best-effort: which blank did this order use? Type from the product, style from the configuration."""
    p = product.lower()
    c = config.lower()
    candidates = [b for b in blanks if b["type"].lower().split()[0] in p or p.split()[0] in b["type"].lower()]
    if not candidates:
        return None
    style_words = {
        "linen": "linen", "white": "white", "stripe": "stripe", "striped": "stripe", "plaid": "plaid",
        "gingham": "gingham", "toile": "toile", "seersucker": "seersucker", "pink": "pink", "blue": "blue",
    }
    for word, key in style_words.items():
        if re.search(rf"\b{word}\b", c):
            hits = [b for b in candidates if key in b["style"].lower()]
            if hits:
                # oldest purchase first (FIFO)
                return sorted(hits, key=lambda b: b["purchased_on"] or "")[0]
    return sorted(candidates, key=lambda b: b["purchased_on"] or "")[0] if len(candidates) == 1 else None


def load_orders(ws, blanks: list[dict], products: list[dict]):
    customers, orders, lines, payments = {}, [], [], []
    for r in sheet_rows(ws):
        d = iso(r.get("Order Date"))
        cust = s(r.get("Customer / Gift"))
        if not d or not cust:
            continue
        ref = s(r.get("Order #")).rstrip(".0") if not isinstance(r.get("Order #"), (int, float)) else str(int(r.get("Order #")))
        key = f"order|ref|{ref.lower()}" if ref else f"order|{cust.lower()}|{d}"
        oid = imp_id(key)
        cid = customers.setdefault(cust.lower(), {"id": imp_id(f"customer|{cust.lower()}"), "name": cust, "contact": "", "notes": "", "created_at": NOW})["id"]

        product = s(r.get("Product"))
        config = s(r.get("Configuration"))
        qty = int(r.get("Qty of Finished Product") or 1)
        revenue = cents(r.get("Total Revenue ($)")) or cents(r.get("Selling Price Before Add-ons ($)"))
        unit_price = round(revenue / qty) if qty else revenue

        completed = iso(r.get("Completion Date"))
        delivered = iso(r.get("Delivery Date"))
        paid_on = iso(r.get("Paid Date"))
        method_raw = s(r.get("Payment Method")).lower()
        method = "venmo" if "venmo" in method_raw else "cash" if "cash" in method_raw else "other"
        gift = s(r.get("Gift / Sample?")).lower() in ("yes", "y", "true", "x")
        status = "delivered" if delivered else "done" if completed else "confirmed"

        notes = []
        deliv = s(r.get("Delivery Method"))
        if deliv:
            notes.append(f"Delivery: {deliv}")
        if gift:
            notes.append("Gift / sample")
        if completed and not delivered:
            notes.append(f"Finished {completed}")

        blank = pick_blank(product, config, blanks)
        unit_cost = cents(r.get("Blank Cost / Unit ($)")) or (blank["unit_cost"] if blank else 0)
        prod = next((p for p in products if p["name"].lower() == product.lower()), None)

        orders.append({
            "id": oid, "customer_id": cid, "customer_name": cust, "status": status, "ordered_on": d,
            "due_on": None, "notes": "; ".join(notes), "created_at": NOW, "updated_at": NOW,
        })
        lines.append({
            "id": imp_id(f"{key}|line|{len([l for l in lines if l['order_id'] == oid])}"),
            "order_id": oid,
            "description": f"{product} - {config}" if config else product,
            "qty": qty, "unit_price": unit_price, "unit_cost": unit_cost, "stitches": None,
            "position": len([l for l in lines if l["order_id"] == oid]),
            "blank_id": blank["id"] if blank else None,
            "product_id": prod["id"] if prod else None,
        })
        if paid_on and revenue > 0:
            payments.append({
                "id": imp_id(f"{key}|payment"), "order_id": oid, "amount": revenue, "method": method,
                "received_on": paid_on, "note": "imported",
            })
    return list(customers.values()), orders, lines, payments


def scan_designs(root: str) -> list[tuple[str, int, str]]:
    """(name, stitches, note) for each design file on the drive, skipping fonts/reference."""
    found = []
    if not root or not os.path.isdir(root):
        return found
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith("__MACOSX")]
        for fn in filenames:
            ext = os.path.splitext(fn)[1].lower()
            if ext not in (".exp", ".pes") or fn.startswith("._"):
                continue
            path = os.path.join(dirpath, fn)
            res = parse_exp(path) if ext == ".exp" else parse_pes(path)
            if not res or not res[0]:
                continue
            rel = os.path.relpath(path, root)
            top = rel.split(os.sep)[0]
            found.append((top, res[0], rel.replace(os.sep, "/")))
    return found


def load_designs(ws, drive_designs) -> list[dict]:
    # group drive files by their top folder: {folder: [(stitches, rel)]}
    by_folder: dict[str, list] = {}
    for top, st, rel in drive_designs:
        by_folder.setdefault(top, []).append((st, rel))

    ALIASES = {
        "duck": "Mallard",
        "dog with duck": "Dog with Mallard Outline",
        "blue/green wallpaper pumpkin": "Blue & Green Pumpkin",
        "basket-weave pumpkin with bow": "Wicker Pumpkin with Bow",
        "cotton stalk": "Cotton Stalk",
        "magnolia": "Magnolia",
        "deer": "Deer Sketch",
        "turkey outline": "Turkey Outline",
    }
    STOP = {"with", "and", "the", "bow", "in", "outline", "sketch"}

    def match_folder(name: str):
        alias = ALIASES.get(name.strip().lower())
        if alias and alias in by_folder:
            return alias
        words = [w for w in re.sub(r"[^a-z ]", " ", name.lower()).split() if len(w) > 2 and w not in STOP]
        for folder in by_folder:
            fw = set(re.sub(r"[^a-z ]", " ", folder.lower()).split())
            if words and all(w in fw for w in words):
                return folder
        return None

    def size_label(rel: str) -> str:
        parts = rel.split("/")
        sub = parts[1] if len(parts) > 2 else ""
        m = re.search(r"(\d(?:\.\d)?)\s*(?:in|inch|inches)?", sub) if sub else None
        if m:
            return f"{m.group(1)} in"
        stem = os.path.splitext(os.path.basename(rel))[0]
        m = re.search(r"(\d(?:\.\d)?)\s*(?:in|inch|inches)", stem, re.I) or re.search(r"(\d\.\d)", stem)
        return f"{m.group(1)} in" if m else stem

    out = []
    for i, r in enumerate(sheet_rows(ws)):
        name = s(r.get("Design Name"))
        if not name:
            continue
        coll = s(r.get("Collection"))
        vendor = s(r.get("Vendor"))
        folder = match_folder(name)
        stitches = 0
        notes = [x for x in (coll, vendor) if x]
        if folder:
            seen: dict[str, int] = {}
            for st, rel in sorted(by_folder[folder]):
                lbl = size_label(rel)
                if lbl not in seen:
                    seen[lbl] = st
            sizes = sorted(seen.items(), key=lambda kv: kv[1])
            stitches = sizes[0][1]
            if len(sizes) > 1:
                notes.append("sizes: " + ", ".join(f"{lbl} {st:,} st" for lbl, st in sizes[:8]))
            else:
                notes.append(f"{sizes[0][0]} {sizes[0][1]:,} st")
        out.append({
            "id": imp_id(f"design|{name.lower()}"),
            "name": name,
            "stitches": stitches,
            "default_price": 0,
            "default_cost": 0,
            "notes": " · ".join(notes),
            "last_used_at": None,
            "created_at": NOW,
        })
    return out


# ---------------------------------------------------------------------------
# apply

def run_sql(statements: list[str], dry: bool):
    """Batch one-line statements through the CLI, staying under the arg-length limit."""
    batch, size, n = [], 0, 0
    def flush():
        nonlocal batch, size, n
        if not batch:
            return
        sql = " ".join(batch)
        if dry:
            print(f"  [dry-run] batch of {len(batch)} statements, {len(sql)} chars")
        else:
            r = subprocess.run(["npx.cmd", "supabase", "db", "query", "--linked", "--", sql], capture_output=True, text=True)
            out = r.stdout + r.stderr
            if r.returncode != 0 or '"Error"' in out:
                print(out[-1500:], file=sys.stderr)
                sys.exit(f"batch failed at statement {n}")
        n += len(batch)
        batch, size = [], 0
    for st in statements:
        if size + len(st) > 2500:
            flush()
        batch.append(st)
        size += len(st) + 1
    flush()
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--designs", default="", help="folder of purchased design files (stitch counts)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.workbook, data_only=True)
    cats: dict[str, str] = {
        "Blanks & garments": "cat-blanks", "Thread & stabilizer": "cat-thread", "Machine & equipment": "cat-machine",
        "Software & designs": "cat-software", "Shipping & packaging": "cat-shipping", "Marketing": "cat-marketing",
        "Mileage & travel": "cat-mileage", "Other": "cat-other",
    }
    known = set(cats)

    expenses = load_expenses(wb["Startup & Fixed Costs"], cats)
    blanks, blank_expenses = load_blanks(wb["Blank Inventory"], cats, expenses)
    products = load_products(wb["Product Pricing"])
    customers, orders, lines, payments = load_orders(wb["Orders"], blanks, products)
    designs = load_designs(wb["Design Library"], scan_designs(a.designs))

    new_cats = [{"id": cid, "name": name, "position": 100 + i} for i, (name, cid) in enumerate(cats.items()) if name not in known]

    print(f"expenses {len(expenses)} (+{len(blank_expenses)} blank purchases) | blanks {len(blanks)} | products {len(products)} | "
          f"customers {len(customers)} | orders {len(orders)} | lines {len(lines)} | payments {len(payments)} | designs {len(designs)} | new categories {[c['name'] for c in new_cats]}")
    print(f"startup total: ${sum(e['amount'] for e in expenses + blank_expenses if e['is_startup'])/100:,.2f}   "
          f"revenue: ${sum(l['unit_price']*l['qty'] for l in lines)/100:,.2f}   paid: ${sum(p['amount'] for p in payments)/100:,.2f}")
    linked = sum(1 for l in lines if l["blank_id"])
    print(f"order lines linked to a blank: {linked}/{len(lines)}; with a product: {sum(1 for l in lines if l['product_id'])}")
    for d in designs:
        print(f"  design: {d['name']:<32} {d['stitches']:>6,} st  {d['notes'][:80]}")

    stmts = []
    stmts += [upsert("expense_category", c) for c in new_cats]
    stmts += [upsert("customer", c) for c in customers]
    stmts += [upsert("order", o) for o in orders]
    stmts += [f"delete from public.order_line where order_id in ({', '.join(q(o['id']) for o in orders)});"] if orders else []
    stmts += [upsert("order_line", l) for l in lines]
    stmts += [upsert("payment", p) for p in payments]
    stmts += [upsert("expense", e) for e in expenses + blank_expenses]
    stmts += [upsert("blank", b) for b in blanks]
    stmts += [upsert("product", p) for p in products]
    stmts += [upsert("design", d) for d in designs]
    n = run_sql(stmts, a.dry_run)
    print(f"{'would apply' if a.dry_run else 'applied'} {len(stmts)} statements ({n} sent)")


if __name__ == "__main__":
    main()
